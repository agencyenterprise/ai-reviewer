"""Tests for checking a quote against the document before it is offered.

The add-in finds a quote in Word's paragraphs; the check reads the agent's markdown.
What matters is that the two agree: markdown syntax is not text, a table cell is a
paragraph, and quotes, dashes and spacing compare as the add-in compares them.
"""

import asyncio
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from lib.services.converters.base import convert_to_markdown
from lib.services.microsoft.word.quote_check import normalize, paragraph_texts, place

MARKDOWN = """\
# Energy *Considerations*

The 2.1-GW figure is inferred from a **generic** conversion.

| Site | Capacity |
| --- | --- |
| Rockport | Inadequate for large scale development |
| Fermi | Inadequate for large scale development |

See [the appendix](#appendix) for site\\_specific data – BTM.
"""


def test_markdown_syntax_is_not_text() -> None:
    paragraphs = paragraph_texts(MARKDOWN)

    assert "Energy Considerations" in paragraphs
    assert "The 2.1-GW figure is inferred from a generic conversion." in paragraphs
    assert "See the appendix for site_specific data - BTM." in paragraphs


def test_each_table_cell_is_its_own_paragraph() -> None:
    paragraphs = paragraph_texts(MARKDOWN)

    assert "Rockport" in paragraphs
    assert "---" not in " ".join(paragraphs)


def test_a_quote_in_one_paragraph_is_found() -> None:
    assert place("inferred from a generic conversion", paragraph_texts(MARKDOWN)) == (
        "found"
    )


def test_a_quote_in_two_paragraphs_is_ambiguous() -> None:
    paragraphs = paragraph_texts(MARKDOWN)

    assert place("Inadequate for large scale development", paragraphs) == "ambiguous"


def test_a_quote_twice_in_one_paragraph_is_ambiguous() -> None:
    """The add-in would change the first, whichever one was meant."""

    assert place("the rate", ["the rate rose, and the rate fell"]) == "ambiguous"


def test_a_quote_that_is_not_there_is_not_found() -> None:
    assert place("a figure nobody wrote", paragraph_texts(MARKDOWN)) == "not_found"
    assert place("   ", paragraph_texts(MARKDOWN)) == "not_found"


def test_dashes_quotes_and_spacing_compare_as_the_add_in_compares_them() -> None:
    paragraphs = paragraph_texts(MARKDOWN)

    assert place("data — BTM", paragraphs) == "found"
    assert normalize("“Site”  specific – BTM") == '"Site" specific - BTM'


# What markitdown produces for these Word paragraphs: escaped `_` and `*`, a literal `|`
# in prose, and two paragraphs of one table cell joined by two spaces. From review n3.
MARKITDOWN = """\
The variable file\\_name holds the path.

Growth was 2\\*3 percent in the base case.

Options are A | B | C for the respondent.

|  |  |
| --- | --- |
| Line one | first  second para in cell |
"""


@pytest.mark.parametrize(
    "quote",
    ["file_name holds", "2*3 percent", "A | B | C for the respondent"],
)
def test_a_verbatim_quote_with_a_lone_underscore_star_or_pipe_is_found(
    quote: str,
) -> None:
    assert place(quote, paragraph_texts(MARKITDOWN)) == "found"


def test_two_paragraphs_of_one_table_cell_stay_apart() -> None:
    paragraphs = paragraph_texts(MARKITDOWN)

    assert place("first second para", paragraphs) == "not_found"
    assert place("second para in cell", paragraphs) == "found"


def test_a_soft_line_break_does_not_split_a_paragraph() -> None:
    markdown = "Results were mixed\nacross all sites.\n\nNext paragraph."

    assert place("mixed across all", paragraph_texts(markdown)) == "found"


def test_list_items_are_paragraphs_of_their_own() -> None:
    paragraphs = paragraph_texts("* one item\n* two items\n\nAfter.")

    assert paragraphs == ["one item", "two items", "After."]


def test_paired_emphasis_is_markup_but_a_lone_marker_is_text() -> None:
    paragraphs = paragraph_texts("A **bold** claim and *italic* aside, 2*3 and a_b.")

    assert paragraphs == ["A bold claim and italic aside, 2*3 and a_b."]


def test_the_rules_hold_on_markitdowns_own_output(tmp_path: Path) -> None:
    """Built from a real .docx, so a change in markitdown's output shows up here."""

    document = DocxDocument()
    soft = document.add_paragraph("Results were mixed")
    soft.add_run().add_break()
    soft.add_run("across all sites.")
    document.add_paragraph("The variable file_name holds 2*3 and A | B.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Line one"
    cell = table.cell(0, 1)
    cell.text = "first"
    cell.add_paragraph("second para in cell")
    path = tmp_path / "probe.docx"
    document.save(str(path))

    markdown = asyncio.run(convert_to_markdown(str(path), converter="markitdown"))
    paragraphs = paragraph_texts(markdown)

    assert place("mixed across all", paragraphs) == "found"
    assert place("file_name holds 2*3 and A | B", paragraphs) == "found"
    assert place("second para in cell", paragraphs) == "found"
    assert place("first second para", paragraphs) == "not_found"
