"""The .docx the Word pane sends, read for the chat agent."""

import io
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from docx import Document as NewDocument
from docx.shared import Inches
from PIL import Image

from lib.config.env import config
from lib.services.microsoft.word.pane_document import NotADocument, from_docx


def a_docx(*paragraphs: str) -> bytes:
    document = NewDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.mark.asyncio
async def test_a_document_is_read_as_markdown() -> None:
    read = await from_docx(
        a_docx("The rate rose by 4%.", "Funding is discussed below."),
        name="Draft.docx",
        url="https://x/Draft.docx",
    )

    assert "The rate rose by 4%." in read.markdown
    assert "Funding is discussed below." in read.markdown
    assert read.name == "Draft.docx" and read.size_bytes > 0
    assert read.images == []


@pytest.mark.asyncio
async def test_a_file_that_is_not_a_zip_is_refused() -> None:
    with pytest.raises(NotADocument, match="not a .docx"):
        await from_docx(b"<x/>", name="x.docx", url="")


@pytest.mark.asyncio
async def test_a_zip_without_a_document_is_refused() -> None:
    out = io.BytesIO()
    with ZipFile(out, "w") as archive:
        archive.writestr("readme.txt", "hello")
    with pytest.raises(NotADocument, match="word/document.xml"):
        await from_docx(out.getvalue(), name="x.docx", url="")


@pytest.mark.asyncio
async def test_a_picture_is_saved_and_referenced_for_view_image(tmp_path: Path) -> None:
    picture = io.BytesIO()
    Image.new("RGB", (40, 20), "red").save(picture, format="PNG")
    document = NewDocument()
    document.add_paragraph("Figure 1. A red box")
    document.add_picture(io.BytesIO(picture.getvalue()), width=Inches(1))
    out = io.BytesIO()
    document.save(out)

    with patch.object(config, "FILE_UPLOADS_MOUNT_PATH", str(tmp_path)):
        read = await from_docx(out.getvalue(), name="x.docx", url="")

    (image,) = read.images
    assert f"draftdetective://{image.id}" in read.markdown
    assert "base64" not in read.markdown
    assert image.mime_type == "image/png" and Path(image.image_path).is_file()
