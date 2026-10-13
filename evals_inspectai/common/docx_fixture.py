"""Build a reviewer's marked-up Word copy from readable fixture text.

The fixtures stay markdown so they can be read and diffed; the Word file a
reviewer would have returned, with their comments in the margin, is made from
that markdown when the sample is uploaded. Shared by the review-assistant
suites that hand the Review Assistant a reviewer who only left comments.
"""

import re
import tempfile
from pathlib import Path

from docx import Document
from docx_editor import Document as EditorDocument
from pydantic import BaseModel, Field

from evals_inspectai.common.loaders import resolve_input


class FixtureComment(BaseModel):
    anchor: str
    text: str


class MarkedUpCopy(BaseModel):
    """A dataset record's marked-up copy: which draft, whose comments, and on what."""

    file_name: str
    marked_up_from: str = Field(description="`file://` reference to the draft's markdown")
    author: str
    comments: list[FixtureComment]

    def content(self) -> bytes:
        return marked_up_docx(resolve_input(self.marked_up_from), self.author, self.comments)


def with_byline(text: str, byline: str) -> str:
    """The memo with ``byline`` under its title line, where a reviewer signs it."""
    title, _, rest = text.partition("\n")
    return f"{title}\n\n{byline}\n{rest}"


def _markdown_to_docx(markdown: str, path: Path) -> None:
    """One paragraph per markdown block, headings as headings. Enough for the
    agent to read the text and for comments to anchor on it."""
    document = Document()
    for block in re.split(r"\n\s*\n", markdown):
        text = " ".join(line.strip() for line in block.strip().splitlines())
        if not text:
            continue
        heading = re.match(r"^(#+)\s+(.*)", text)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 3))
        else:
            document.add_paragraph(text.replace("**", ""))
    document.save(str(path))


def marked_up_docx(markdown: str, author: str, comments: list[FixtureComment]) -> bytes:
    """The markdown as a Word file carrying ``author``'s comments."""
    with tempfile.TemporaryDirectory() as workdir:
        path = Path(workdir) / "marked-up.docx"
        _markdown_to_docx(markdown, path)
        doc = EditorDocument.open(path, author=author, workspace_dir=str(Path(workdir) / "ws"))
        try:
            for comment in comments:
                doc.add_comment(comment.anchor, comment.text)
            doc.save()
        finally:
            doc.close()
        return path.read_bytes()
